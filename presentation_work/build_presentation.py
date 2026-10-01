from pathlib import Path
from math import sin, pi

from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
ASSETS = ROOT / "presentation_work" / "assets"
OUT.mkdir(exist_ok=True)
ASSETS.mkdir(exist_ok=True)

W, H = 13.333, 7.5
NAVY = "0B101A"
SURFACE = "121A28"
SURFACE_2 = "172235"
MINT = "65E6C4"
GREEN = "8CFB65"
PURPLE = "8B7CF6"
WHITE = "F6FAFF"
MUTED = "9EABC0"
BORDER = "29364B"
FONT = "Pretendard"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def make_background(path: Path):
    im = Image.new("RGB", (1600, 900), (9, 14, 24))
    glow = Image.new("RGBA", im.size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(glow)
    draw.ellipse((-260, -300, 760, 720), fill=(50, 210, 190, 72))
    draw.ellipse((950, -280, 1860, 610), fill=(125, 95, 255, 64))
    draw.ellipse((920, 560, 1650, 1180), fill=(30, 110, 150, 34))
    glow = glow.filter(ImageFilter.GaussianBlur(110))
    im = Image.alpha_composite(im.convert("RGBA"), glow)
    grid = Image.new("RGBA", im.size, (0, 0, 0, 0))
    gd = ImageDraw.Draw(grid)
    for x in range(0, 1600, 52):
        gd.line((x, 0, x, 900), fill=(255, 255, 255, 8), width=1)
    for y in range(0, 900, 52):
        gd.line((0, y, 1600, y), fill=(255, 255, 255, 8), width=1)
    Image.alpha_composite(im, grid).convert("RGB").save(path, quality=94)


BG = ASSETS / "background.jpg"
make_background(BG)


def add_bg(slide, index=None):
    slide.shapes.add_picture(str(BG), 0, 0, Inches(W), Inches(H))
    if index is not None:
        add_text(slide, f"{index:02d}", 12.35, 7.08, .45, .2, 8, MUTED, bold=True, align=PP_ALIGN.RIGHT)


def add_text(slide, text, x, y, w, h, size=18, color=WHITE, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, margin=.02):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.clear()
    box.text_frame.margin_left = Inches(margin)
    box.text_frame.margin_right = Inches(margin)
    box.text_frame.margin_top = Inches(margin)
    box.text_frame.margin_bottom = Inches(margin)
    box.text_frame.vertical_anchor = valign
    p = box.text_frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_rect(slide, x, y, w, h, fill=SURFACE, line=BORDER, radius=True, transparency=0):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.fill.transparency = transparency
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(.8)
    return shape


def add_chip(slide, text, x, y, w=None, color=MINT):
    width = w or max(.8, len(text) * .115 + .38)
    shape = add_rect(slide, x, y, width, .34, SURFACE_2, color)
    add_text(slide, text, x, y, width, .34, 9, color, True, PP_ALIGN.CENTER)
    return shape


def add_title(slide, eyebrow, title, subtitle=None, index=None):
    add_bg(slide, index)
    add_text(slide, eyebrow.upper(), .62, .43, 4.8, .28, 9, MINT, True)
    add_text(slide, title, .62, .76, 11.7, .68, 30, WHITE, True)
    if subtitle:
        add_text(slide, subtitle, .64, 1.45, 11.2, .42, 12, MUTED)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.62), Inches(1.96), Inches(12.05), Inches(.015))
    line.fill.solid()
    line.fill.fore_color.rgb = rgb(BORDER)
    line.line.fill.background()


def add_card(slide, x, y, w, h, kicker, title, body, accent=MINT, number=None):
    add_rect(slide, x, y, w, h, SURFACE, BORDER)
    if number:
        add_text(slide, number, x + .25, y + .25, .55, .38, 18, accent, True)
        tx = x + .86
    else:
        tx = x + .28
    add_text(slide, kicker.upper(), tx, y + .24, w - (tx - x) - .2, .24, 8, accent, True)
    add_text(slide, title, x + .28, y + .65, w - .56, .42, 16, WHITE, True)
    add_text(slide, body, x + .28, y + 1.17, w - .56, h - 1.4, 10, MUTED, False, valign=MSO_ANCHOR.TOP)


def add_wave(slide, x, y, w, h, bars=38, accent=MINT):
    bw = w / bars * .48
    gap = w / bars
    for i in range(bars):
        value = .23 + .68 * abs(sin((i + 2) * 1.61) * sin((i + 4) * .47))
        bar_h = h * value
        bar = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x + i * gap),
            Inches(y + (h - bar_h) / 2),
            Inches(bw),
            Inches(bar_h),
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = rgb(accent if i % 5 else GREEN)
        bar.fill.transparency = 12
        bar.line.fill.background()


def add_browser(slide, x, y, w, h, title, mode="composer"):
    add_rect(slide, x, y, w, h, "0D1420", "33435A")
    top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(.36))
    top.fill.solid(); top.fill.fore_color.rgb = rgb("101826"); top.line.fill.background()
    for i, c in enumerate(("FF6B6B", "FFD166", "65E6C4")):
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + .18 + i * .2), Inches(y + .12), Inches(.08), Inches(.08))
        dot.fill.solid(); dot.fill.fore_color.rgb = rgb(c); dot.line.fill.background()
    add_text(slide, title, x + .8, y + .05, w - 1.6, .25, 8, MUTED, False, PP_ALIGN.CENTER)

    if mode == "composer":
        add_text(slide, "멜로디", x + .18, y + .52, .7, .25, 8, MINT, True)
        for r in range(8):
            yy = y + .9 + r * ((h - 1.35) / 8)
            key = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + .18), Inches(yy), Inches(.48), Inches(.25))
            key.fill.solid(); key.fill.fore_color.rgb = rgb("E8EEF6" if r % 2 else "10151E"); key.line.color.rgb = rgb(BORDER)
            for c in range(10):
                cell = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + .72 + c * ((w - 1) / 10)), Inches(yy), Inches((w - 1) / 10), Inches(.25))
                cell.fill.solid(); cell.fill.fore_color.rgb = rgb("171F2C"); cell.line.color.rgb = rgb("253247")
        notes = [(2,1,3),(5,3,2),(7,5,2),(4,7,3)]
        for row,col,length in notes:
            yy = y + .9 + row * ((h - 1.35) / 8)
            xx = x + .72 + col * ((w - 1) / 10)
            note = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(xx), Inches(yy+.03), Inches(length*((w-1)/10)-.03), Inches(.19))
            note.fill.solid(); note.fill.fore_color.rgb = rgb(GREEN if row != 7 else MINT); note.line.fill.background()
        add_wave(slide, x + .75, y + h - .34, w - 1.05, .18, 32)
    elif mode == "community":
        for i in range(3):
            yy = y + .62 + i * 1.15
            add_rect(slide, x + .22, yy, w - .44, .9, SURFACE_2, BORDER)
            add_text(slide, f"{i+1:02d}", x + .4, yy + .16, .38, .28, 14, GREEN, True)
            add_text(slide, ["야경 루프", "감성 브리지", "시티팝 리프"][i], x + .88, yy + .12, 1.4, .3, 10, WHITE, True)
            add_wave(slide, x + 2.35, yy + .22, w - 3.15, .3, 22)
    else:
        add_text(slide, "LIVE CAMERA", x + .3, y + .58, 2, .25, 9, MINT, True)
        stage = add_rect(slide, x + .25, y + .95, w - .5, h - 1.25, "081019", "344A5C")
        stage.fill.transparency = 8
        add_text(slide, "손동작으로\n연주 · 녹화", x + .5, y + 1.35, w - 1, 1.2, 22, WHITE, True, PP_ALIGN.CENTER)
        add_chip(slide, "R 녹화 시작", x + .7, y + h - .7, 1.15, GREEN)
        add_chip(slide, "S 저장", x + 1.95, y + h - .7, .9, MINT)


prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]

# 1. Cover
slide = prs.slides.add_slide(blank)
add_bg(slide, 1)
add_chip(slide, "GRADUATION PROJECT · WEB MUSIC WORKSPACE", .7, .62, 3.55)
add_text(slide, "아이디어가\n곡이 되는 순간", .72, 1.28, 6.4, 1.75, 38, WHITE, True, valign=MSO_ANCHOR.TOP)
add_text(slide, "Song Project", .74, 3.22, 4.8, .62, 25, MINT, True)
add_text(slide, "작곡 · 협업 · 공유 · 에어 악기를 하나의 웹 작업실에서", .76, 3.88, 5.65, .62, 14, MUTED)
add_rect(slide, 7.22, .72, 5.38, 5.95, SURFACE, "33455C")
add_browser(slide, 7.55, 1.05, 4.72, 4.6, "Song Project / Composer", "composer")
add_wave(slide, 7.9, 5.92, 4.05, .38, 40)
add_text(slide, "2026 졸업작품 발표", .75, 6.82, 3, .25, 10, MUTED)

# 2. Contents
slide = prs.slides.add_slide(blank)
add_title(slide, "Presentation Flow", "발표 순서", "기능을 나열하기보다 실제 사용 흐름대로 보여드립니다.", 2)
items = [
    ("01", "프로젝트 개요", "왜 필요한가"),
    ("02", "개발 환경", "어떻게 만들었나"),
    ("03", "서비스 구조", "사용자는 어디서 무엇을 하나"),
    ("04", "핵심 기능 시연", "작곡 · 공유 · 협업 · 에어 악기"),
    ("05", "기술 구조", "실시간 연결과 데이터 흐름"),
    ("06", "기대효과", "무엇이 달라지는가"),
]
for i, (n, title, desc) in enumerate(items):
    col, row = i % 2, i // 2
    x, y = .72 + col * 6.15, 2.3 + row * 1.35
    add_rect(slide, x, y, 5.68, 1.02, SURFACE, BORDER)
    add_text(slide, n, x + .24, y + .18, .55, .4, 18, GREEN, True)
    add_text(slide, title, x + .92, y + .12, 2.2, .35, 14, WHITE, True)
    add_text(slide, desc, x + .92, y + .5, 3.9, .25, 9, MUTED)

# 3. Overview
slide = prs.slides.add_slide(blank)
add_title(slide, "Project Overview", "작곡의 시작부터 공유까지, 흐름이 끊기지 않게", "초보자의 진입장벽과 협업 도구의 분산 문제를 하나의 웹 서비스로 해결합니다.", 3)
add_card(slide, .7, 2.25, 3.8, 3.75, "Problem 01", "작곡은 어렵다", "전문 DAW는 기능이 많고 복잡합니다.\n처음 곡을 만드는 사용자는 멜로디를 기록하기도 전에 도구 사용법에서 막힙니다.", GREEN)
add_card(slide, 4.77, 2.25, 3.8, 3.75, "Problem 02", "협업은 흩어진다", "메신저, 파일 전송, 화상회의를 오가며 버전이 뒤섞입니다.\n누가 어느 파트를 수정했는지 추적하기 어렵습니다.", MINT)
add_card(slide, 8.84, 2.25, 3.8, 3.75, "Solution", "하나의 음악 작업실", "브라우저에서 작곡하고, 실시간으로 함께 수정하며, 완성 전 스케치도 커뮤니티에 공유합니다.", PURPLE)
add_text(slide, "MIDI 감각의 편집  ×  실시간 협업  ×  커뮤니티 피드백", .75, 6.45, 8.5, .38, 15, WHITE, True)

# 4. Goal
slide = prs.slides.add_slide(blank)
add_title(slide, "Core Value", "Song Project가 만드는 3가지 변화", None, 4)
goals = [
    ("01", "Barrier-Free Creation", "설치 없이 브라우저에서 바로 시작", "작곡 화면에 들어가 음을 찍고 즉시 재생합니다.", GREEN),
    ("02", "Interactive Co-Creation", "같은 곡을 실시간으로 공동 편집", "파트별 모집과 채팅으로 협업 진입을 줄입니다.", MINT),
    ("03", "From Sketch to Community", "스케치 단계부터 공유하고 발전", "파형, Song DNA, 댓글 피드백이 다음 작업으로 이어집니다.", PURPLE),
]
for i, (n, en, title, body, accent) in enumerate(goals):
    x = .72 + i * 4.17
    add_card(slide, x, 2.2, 3.82, 3.85, en, title, body, accent, n)
    add_wave(slide, x + .35, 5.42, 3.1, .34, 24, accent)
add_text(slide, "핵심은 ‘기능 수’가 아니라, 작업이 다음 행동으로 자연스럽게 이어지는 경험입니다.", .75, 6.45, 11.8, .38, 13, MUTED)

# 5. Development environment
slide = prs.slides.add_slide(blank)
add_title(slide, "Development Environment", "개발 환경", "빠른 상호작용, 실시간 데이터, 브라우저 오디오를 중심으로 구성했습니다.", 5)
tech = [
    ("FRONTEND", "React 19\nTypeScript\nVite", GREEN),
    ("AUDIO", "Tone.js\nWeb Audio API\nMediaRecorder", MINT),
    ("VISION", "MediaPipe\nCamera Hands\nCanvas", PURPLE),
    ("DATA", "Firebase\nNode.js\nSQLite", GREEN),
    ("REALTIME", "WebSocket\nWebRTC\nConflict Lock", MINT),
]
for i, (label, body, accent) in enumerate(tech):
    x = .72 + i * 2.48
    add_rect(slide, x, 2.35, 2.18, 3.15, SURFACE, BORDER)
    add_chip(slide, label, x + .22, 2.6, 1.45, accent)
    add_text(slide, body, x + .24, 3.25, 1.7, 1.3, 15, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, ["UI / 상태 관리", "재생 / 녹음", "손동작 인식", "저장 / 인증", "공동 편집"][i], x + .22, 4.95, 1.74, .25, 9, MUTED, False, PP_ALIGN.CENTER)
add_text(slide, "웹 표준 기술을 조합해 별도 프로그램 설치 없이 음악 제작 환경을 제공합니다.", .74, 6.18, 10.8, .45, 13, WHITE, True)

# 6. Service map
slide = prs.slides.add_slide(blank)
add_title(slide, "Service Map", "서비스 기능도", "사용자는 작곡을 중심으로 라이브러리·협업·커뮤니티를 순환합니다.", 6)
center = add_rect(slide, 5.13, 3.02, 3.05, 1.18, "173126", MINT)
add_text(slide, "COMPOSER", 5.13, 3.13, 3.05, .28, 10, MINT, True, PP_ALIGN.CENTER)
add_text(slide, "곡 스케치 작업실", 5.13, 3.45, 3.05, .36, 18, WHITE, True, PP_ALIGN.CENTER)
nodes = [
    (.75, 2.2, "MAIN", "추천곡 · HOT 5"),
    (.75, 4.45, "LIBRARY", "저장 · 불러오기"),
    (9.0, 2.2, "COMMUNITY", "공유곡 · 게시판"),
    (9.0, 4.45, "COLLAB", "파트 모집 · 채팅"),
]
for x, y, k, t in nodes:
    add_rect(slide, x, y, 3.55, 1.3, SURFACE, BORDER)
    add_text(slide, k, x + .24, y + .19, 1.5, .25, 9, GREEN, True)
    add_text(slide, t, x + .24, y + .55, 2.9, .34, 15, WHITE, True)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.3 if x < 5 else 8.18), Inches(y + .63), Inches(.83), Inches(.025))
    line.fill.solid(); line.fill.fore_color.rgb = rgb(MINT); line.line.fill.background()
add_chip(slide, "VIDEO OVERLAY", 5.03, 2.38, 1.45, PURPLE)
add_chip(slide, "AIR INSTRUMENT", 6.66, 2.38, 1.55, GREEN)
add_chip(slide, "SONG DNA", 5.49, 4.53, 1.15, MINT)
add_chip(slide, "MP3 / WEBM", 6.82, 4.53, 1.18, PURPLE)

# 7. Composer demo
slide = prs.slides.add_slide(blank)
add_title(slide, "Feature Demo 01", "작곡 화면 — 찍고, 듣고, 바로 수정", "설명보다 시연이 중심이 되는 핵심 화면입니다.", 7)
add_browser(slide, .7, 2.2, 7.5, 4.48, "Song Project / Composer", "composer")
features = [
    ("01", "피아노 롤 입력", "음정과 길이를 눈으로 확인"),
    ("02", "악기별 트랙", "멜로디·기타·드럼·베이스"),
    ("03", "템포·반복 재생", "마디 단위로 빠르게 확인"),
    ("04", "영상 오버레이", "영상에 맞춰 음악을 제작"),
]
for i, (n, title, body) in enumerate(features):
    y = 2.22 + i * 1.08
    add_rect(slide, 8.55, y, 4.05, .84, SURFACE, BORDER)
    add_text(slide, n, 8.78, y + .16, .38, .3, 12, GREEN, True)
    add_text(slide, title, 9.3, y + .1, 2.9, .28, 12, WHITE, True)
    add_text(slide, body, 9.3, y + .43, 2.9, .2, 8, MUTED)
add_chip(slide, "DEMO POINT", 8.58, 6.15, 1.1, MINT)
add_text(slide, "음을 입력 → 재생 → 악기 음량 조절", 9.82, 6.11, 2.7, .35, 10, WHITE, True)

# 8. Community / sharing
slide = prs.slides.add_slide(blank)
add_title(slide, "Feature Demo 02", "공유곡이 다음 작업으로 이어지는 구조", "감상에서 끝나지 않고, 저장·피드백·협업 모집으로 연결됩니다.", 8)
add_browser(slide, .72, 2.25, 6.0, 4.15, "Song Project / Music Share", "community")
add_card(slide, 7.05, 2.25, 2.65, 1.8, "01 / LISTEN", "파형 미리보기", "업로드된 곡의 실제 프로젝트 이벤트를 파형으로 표시", GREEN)
add_card(slide, 9.98, 2.25, 2.65, 1.8, "02 / REACT", "댓글 피드백", "마디를 지정해 구체적인 의견을 남김", MINT)
add_card(slide, 7.05, 4.32, 2.65, 1.8, "03 / REUSE", "내 작업실로", "공유곡을 기반으로 새 스케치를 시작", PURPLE)
add_card(slide, 9.98, 4.32, 2.65, 1.8, "04 / CONNECT", "파트 모집", "곡의 장르와 필요한 파트를 협업으로 연결", GREEN)

# 9. Realtime collaboration
slide = prs.slides.add_slide(blank)
add_title(slide, "Feature Demo 03", "실시간 공동 작곡", "같은 프로젝트를 보면서 편집·채팅·파트 상태를 공유합니다.", 9)
steps = [
    ("1", "방 생성", "작업 프로젝트와 필요한 파트를 등록"),
    ("2", "파트 참여", "세션 모집에서 역할을 선택"),
    ("3", "실시간 편집", "WebSocket으로 작업 상태 동기화"),
    ("4", "충돌 방지", "편집 잠금과 버전 정보로 덮어쓰기 방지"),
]
for i, (n, title, body) in enumerate(steps):
    x = .72 + i * 3.08
    add_rect(slide, x, 2.55, 2.72, 2.45, SURFACE, BORDER)
    add_text(slide, n, x + .2, 2.77, .45, .45, 20, GREEN if i < 3 else PURPLE, True, PP_ALIGN.CENTER)
    add_text(slide, title, x + .25, 3.35, 2.2, .34, 14, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, body, x + .32, 3.88, 2.08, .72, 9, MUTED, False, PP_ALIGN.CENTER, MSO_ANCHOR.TOP)
    if i < 3:
        add_text(slide, "→", x + 2.72, 3.55, .36, .3, 19, MINT, True, PP_ALIGN.CENTER)
add_rect(slide, 2.3, 5.48, 8.73, .75, "12251F", "2D6C5A")
add_text(slide, "WebSocket Server  ·  Live Collaboration  ·  Conflict Resolution", 2.3, 5.63, 8.73, .3, 13, MINT, True, PP_ALIGN.CENTER)

# 10. Air instrument
slide = prs.slides.add_slide(blank)
add_title(slide, "Feature Demo 04", "카메라가 악기가 되는 에어 악기", "MediaPipe Hands로 손동작을 인식하고, 연주 화면과 소리를 WebM으로 기록합니다.", 10)
add_browser(slide, .72, 2.22, 6.6, 4.35, "Song Project / Air Instrument", "air")
add_card(slide, 7.65, 2.22, 2.35, 1.72, "MODE 01", "에어 기타", "스트럼과 줄 연주 방식", GREEN)
add_card(slide, 10.28, 2.22, 2.35, 1.72, "MODE 02", "에어 드럼", "패드 위치를 손으로 타격", MINT)
add_card(slide, 7.65, 4.25, 2.35, 1.72, "MODE 03", "에어 피아노", "손가락 위치로 음을 연주", PURPLE)
add_card(slide, 10.28, 4.25, 2.35, 1.72, "RECORD", "R 시작 · S 종료", "종료하면 WebM 자동 저장", GREEN)

# 11. Differentiators
slide = prs.slides.add_slide(blank)
add_title(slide, "Differentiation", "Song Project만의 연결 기능", "각 기능이 독립적으로 끝나지 않고 다음 단계로 이어집니다.", 11)
diffs = [
    ("SONG DNA", "곡 스케치의 분위기·멜로디 유형·활용도를 요약", GREEN),
    ("BRING TO STUDIO", "공유곡을 내 작업실로 가져와 이어서 편곡", MINT),
    ("AUTO PART MATCH", "곡 정보와 모집 파트를 연결해 세션 모집 시작", PURPLE),
    ("CAMERA TO MUSIC", "에어 악기 연주를 녹화하고 작곡 경험으로 확장", GREEN),
]
for i, (title, body, accent) in enumerate(diffs):
    col, row = i % 2, i // 2
    x, y = .72 + col * 6.12, 2.25 + row * 2.05
    add_rect(slide, x, y, 5.72, 1.58, SURFACE, BORDER)
    add_chip(slide, title, x + .25, y + .24, 1.55, accent)
    add_text(slide, body, x + .25, y + .76, 5.15, .46, 12, WHITE, True)
add_wave(slide, 2.1, 6.35, 9.15, .38, 70)

# 12. Architecture
slide = prs.slides.add_slide(blank)
add_title(slide, "System Architecture", "데이터와 미디어 흐름", "브라우저의 즉각적인 반응과 서버의 영속성을 분리했습니다.", 12)
layers = [
    ("CLIENT", "React UI · Zustand · Canvas", .82, GREEN),
    ("MEDIA", "Tone.js · Web Audio · MediaRecorder", 3.75, MINT),
    ("REALTIME", "WebSocket · WebRTC · Locks", 6.68, PURPLE),
    ("SERVER", "Firebase · Node.js · SQLite", 9.61, GREEN),
]
for i, (k, body, x, accent) in enumerate(layers):
    add_rect(slide, x, 2.55, 2.55, 2.65, SURFACE, BORDER)
    add_chip(slide, k, x + .25, 2.82, 1.1, accent)
    add_text(slide, body, x + .3, 3.55, 1.95, 1.0, 13, WHITE, True, PP_ALIGN.CENTER)
    if i < 3:
        add_text(slide, "↔", x + 2.55, 3.62, .38, .3, 18, MINT, True, PP_ALIGN.CENTER)
add_text(slide, "입력은 즉시 반응하고 · 협업 상태는 동기화되며 · 프로젝트는 안전하게 저장됩니다.", 1.25, 5.85, 10.8, .5, 14, MUTED, True, PP_ALIGN.CENTER)

# 13. Outcome
slide = prs.slides.add_slide(blank)
add_title(slide, "Expected Outcome", "플랫폼 기대효과", "음악 제작의 첫 진입과 협업의 다음 행동을 동시에 단순화합니다.", 13)
outcomes = [
    ("진입장벽 완화", "설치와 복잡한 설정 없이 작곡 시작", "+"),
    ("협업 속도 향상", "파일 전달 대신 같은 프로젝트를 편집", "↗"),
    ("피드백의 구체화", "파형과 마디 기준으로 의견 교환", "◎"),
    ("창작 생태계 형성", "공유곡이 새 작업과 파트 모집으로 연결", "∞"),
]
for i, (title, body, icon) in enumerate(outcomes):
    x = .75 + i * 3.08
    add_rect(slide, x, 2.42, 2.72, 3.15, SURFACE, BORDER)
    add_text(slide, icon, x + .22, 2.67, 2.28, .65, 28, GREEN if i % 2 == 0 else MINT, True, PP_ALIGN.CENTER)
    add_text(slide, title, x + .25, 3.62, 2.22, .35, 15, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, body, x + .34, 4.28, 2.04, .65, 10, MUTED, False, PP_ALIGN.CENTER, MSO_ANCHOR.TOP)
add_text(slide, "아이디어 → 스케치 → 협업 → 공유 → 다시 창작", 2.45, 6.1, 8.45, .45, 18, MINT, True, PP_ALIGN.CENTER)

# 14. Closing
slide = prs.slides.add_slide(blank)
add_bg(slide, 14)
add_chip(slide, "THANK YOU", 5.53, 1.05, 2.27, MINT)
add_text(slide, "음악을 만드는 과정이\n혼자가 아니도록", 1.25, 2.0, 10.83, 1.5, 35, WHITE, True, PP_ALIGN.CENTER)
add_text(slide, "Song Project", 4.2, 3.86, 4.93, .62, 25, GREEN, True, PP_ALIGN.CENTER)
add_wave(slide, 3.2, 4.9, 6.95, .52, 56)
add_text(slide, "Q & A", 5.45, 5.82, 2.45, .58, 24, WHITE, True, PP_ALIGN.CENTER)
add_text(slide, "작곡 · 협업 · 공유 · 에어 악기", 4.3, 6.52, 4.73, .3, 10, MUTED, True, PP_ALIGN.CENTER)

for slide in prs.slides:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(NAVY)

output = OUT / "Song_Project_발표자료_수정본.pptx"
prs.save(output)
print(output)

