from io import BytesIO
from pathlib import Path
import textwrap

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR


ROOT = Path(__file__).resolve().parents[1]
PPTX = ROOT / "deliverables" / "Song_Project_발표자료_수정본.pptx"
OUT = ROOT / "deliverables"
PREVIEW = OUT / "preview"
PREVIEW.mkdir(exist_ok=True)

FONT_REGULAR = "C:/Windows/Fonts/NotoSansKR-VF.ttf"
FONT_BOLD = "C:/Windows/Fonts/malgunbd.ttf"
SCALE = 120
PX_W, PX_H = 1600, 900


def emu_to_px(value, total_emu, total_px):
    return round(value / total_emu * total_px)


def get_rgb(color_format, fallback=(20, 28, 42)):
    try:
        value = color_format.rgb
        if value is None:
            return fallback
        return tuple(value)
    except Exception:
        return fallback


def font(size, bold=False):
    path = FONT_BOLD if bold else FONT_REGULAR
    return ImageFont.truetype(path, max(8, round(size * 1.34)))


def fit_lines(draw, text, font_obj, max_width):
    lines = []
    for paragraph in str(text).splitlines() or [""]:
        if not paragraph:
            lines.append("")
            continue
        current = ""
        for char in paragraph:
            test = current + char
            if draw.textlength(test, font=font_obj) <= max_width or not current:
                current = test
            else:
                lines.append(current)
                current = char
        lines.append(current)
    return lines


def render_text(draw, shape, box):
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    x, y, w, h = box
    ml = emu_to_px(tf.margin_left or 0, prs.slide_width, PX_W)
    mr = emu_to_px(tf.margin_right or 0, prs.slide_width, PX_W)
    mt = emu_to_px(tf.margin_top or 0, prs.slide_height, PX_H)
    mb = emu_to_px(tf.margin_bottom or 0, prs.slide_height, PX_H)
    content = []
    for p in tf.paragraphs:
        text = "".join(run.text for run in p.runs) if p.runs else p.text
        if not text:
            continue
        run = p.runs[0] if p.runs else None
        size = float((run.font.size.pt if run and run.font.size else 12))
        bold = bool(run.font.bold) if run else False
        color = get_rgb(run.font.color, (246, 250, 255)) if run else (246, 250, 255)
        fnt = font(size, bold)
        lines = fit_lines(draw, text, fnt, max(1, w - ml - mr))
        align = p.alignment
        line_height = round(size * 1.75)
        for line in lines:
            content.append((line, fnt, color, align, line_height))
    total_h = sum(item[4] for item in content)
    if tf.vertical_anchor == MSO_ANCHOR.MIDDLE:
        cy = y + max(mt, (h - total_h) // 2)
    elif tf.vertical_anchor == MSO_ANCHOR.BOTTOM:
        cy = y + h - mb - total_h
    else:
        cy = y + mt
    for line, fnt, color, align, line_height in content:
        width = draw.textlength(line, font=fnt)
        if align == PP_ALIGN.CENTER:
            cx = x + (w - width) / 2
        elif align == PP_ALIGN.RIGHT:
            cx = x + w - mr - width
        else:
            cx = x + ml
        draw.text((cx, cy), line, font=fnt, fill=color)
        cy += line_height


def render_shape(canvas, draw, shape):
    x = emu_to_px(shape.left, prs.slide_width, PX_W)
    y = emu_to_px(shape.top, prs.slide_height, PX_H)
    w = max(1, emu_to_px(shape.width, prs.slide_width, PX_W))
    h = max(1, emu_to_px(shape.height, prs.slide_height, PX_H))
    box = (x, y, w, h)

    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        try:
            pic = Image.open(BytesIO(shape.image.blob)).convert("RGB")
            pic = pic.resize((w, h), Image.Resampling.LANCZOS)
            canvas.paste(pic, (x, y))
        except Exception:
            pass
        return

    if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
        fill = None
        try:
            if shape.fill.type is not None:
                fill = get_rgb(shape.fill.fore_color)
        except Exception:
            pass
        outline = None
        try:
            outline = get_rgb(shape.line.color, (41, 54, 75))
        except Exception:
            pass
        xy = (x, y, x + w, y + h)
        if shape.auto_shape_type == MSO_SHAPE.OVAL:
            draw.ellipse(xy, fill=fill, outline=outline, width=1)
        elif shape.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
            draw.rounded_rectangle(xy, radius=min(18, h // 3), fill=fill, outline=outline, width=1)
        else:
            draw.rectangle(xy, fill=fill, outline=outline, width=1)
        render_text(draw, shape, box)
        return

    if shape.has_text_frame:
        render_text(draw, shape, box)


prs = Presentation(PPTX)
pages = []
for index, slide in enumerate(prs.slides, 1):
    canvas = Image.new("RGB", (PX_W, PX_H), (9, 14, 24))
    draw = ImageDraw.Draw(canvas)
    for shape in slide.shapes:
        render_shape(canvas, draw, shape)
    page_path = PREVIEW / f"slide-{index:02d}.png"
    canvas.save(page_path)
    pages.append(canvas)

pdf_path = OUT / "Song_Project_발표자료_수정본.pdf"
pages[0].save(pdf_path, save_all=True, append_images=pages[1:], resolution=120)

thumbs = []
for index, page in enumerate(pages, 1):
    thumb = page.copy()
    thumb.thumbnail((480, 270))
    cell = Image.new("RGB", (500, 300), (225, 225, 225))
    cell.paste(thumb, ((500 - thumb.width) // 2, 22))
    ImageDraw.Draw(cell).text((10, 4), str(index), fill=(0, 0, 0), font=font(9, True))
    thumbs.append(cell)
rows = (len(thumbs) + 2) // 3
sheet = Image.new("RGB", (1500, rows * 300), (215, 215, 215))
for i, thumb in enumerate(thumbs):
    sheet.paste(thumb, ((i % 3) * 500, (i // 3) * 300))
sheet.save(OUT / "Song_Project_발표자료_미리보기.jpg", quality=90)

print(pdf_path)
